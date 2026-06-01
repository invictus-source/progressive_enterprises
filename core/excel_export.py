import os
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

import config
from db.manager import get_db
from db.models import Sale, Purchase


HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
SUBHEADER_FILL = PatternFill("solid", fgColor="2563EB")
ALT_FILL = PatternFill("solid", fgColor="EFF6FF")
TOTAL_FILL = PatternFill("solid", fgColor="DBEAFE")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header_cell(ws, row, col, text, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def _data_cell(ws, row, col, value, alt=False, bold=False, number_format=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=9)
    c.fill = ALT_FILL if alt else PatternFill()
    c.alignment = Alignment(vertical="center")
    c.border = BORDER
    if number_format:
        c.number_format = number_format
    return c


def _auto_col_widths(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _out_path(name: str) -> str:
    out_dir = config.EXPORTS_DIR
    os.makedirs(out_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(out_dir, f"{name}_{stamp}.xlsx")
    return path


def export_sales_report(from_date: date, to_date: date) -> str:
    session = get_db()
    try:
        start = datetime.combine(from_date, datetime.min.time())
        end = datetime.combine(to_date, datetime.max.time())
        sales = session.query(Sale).filter(
            Sale.sale_date.between(start, end), Sale.is_cancelled == False
        ).order_by(Sale.sale_date).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = f"{config.COMPANY_NAME} – Sales Report ({from_date} to {to_date})"
        title_cell.font = Font(bold=True, size=13, color="1E3A5F")
        title_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 24

        headers = ["Invoice No.", "Date", "Customer", "Items", "Subtotal", "Discount", "Total GST", "Grand Total"]
        widths = [18, 14, 25, 8, 14, 12, 12, 14]
        for c, (h, w) in enumerate(zip(headers, widths), 1):
            _header_cell(ws, 2, c, h, w)
        ws.row_dimensions[2].height = 20

        total_grand = 0.0
        for r, sale in enumerate(sales, 3):
            alt = r % 2 == 0
            _data_cell(ws, r, 1, sale.invoice_no, alt)
            _data_cell(ws, r, 2, sale.sale_date.strftime("%d-%m-%Y"), alt)
            _data_cell(ws, r, 3, sale.customer.name if sale.customer else "Walk-in", alt)
            _data_cell(ws, r, 4, len(sale.items), alt)
            _data_cell(ws, r, 5, sale.subtotal, alt, number_format='₹#,##0.00')
            _data_cell(ws, r, 6, sale.discount_amount, alt, number_format='₹#,##0.00')
            _data_cell(ws, r, 7, sale.total_gst, alt, number_format='₹#,##0.00')
            _data_cell(ws, r, 8, sale.grand_total, alt, number_format='₹#,##0.00')
            total_grand += sale.grand_total

        tr = len(sales) + 3
        tc = ws.cell(row=tr, column=1, value="TOTAL")
        tc.font = Font(bold=True, size=10, color="1E3A5F")
        tc.fill = TOTAL_FILL
        for col in range(2, 8):
            ws.cell(row=tr, column=col).fill = TOTAL_FILL
        tg = ws.cell(row=tr, column=8, value=total_grand)
        tg.font = Font(bold=True, size=10, color="1E3A5F")
        tg.fill = TOTAL_FILL
        tg.number_format = '₹#,##0.00'
        tg.border = BORDER

        path = _out_path("sales_report")
        wb.save(path)
        return path
    finally:
        session.close()


def export_gst_report(month: int, year: int) -> str:
    session = get_db()
    try:
        start = datetime(year, month, 1)
        end = datetime(year + (month == 12), (month % 12) + 1, 1)
        from calendar import month_name
        period = f"{month_name[month]} {year}"

        sales = session.query(Sale).filter(
            Sale.sale_date >= start, Sale.sale_date < end, Sale.is_cancelled == False
        ).order_by(Sale.sale_date).all()
        purchases = session.query(Purchase).filter(
            Purchase.purchase_date >= start, Purchase.purchase_date < end
        ).order_by(Purchase.purchase_date).all()

        wb = Workbook()

        ws1 = wb.active
        ws1.title = "Outward Supplies"
        ws1.merge_cells("A1:H1")
        ws1["A1"].value = f"GSTR-1 Style – Outward Supplies – {period}"
        ws1["A1"].font = Font(bold=True, size=12, color="1E3A5F")

        heads = ["Invoice", "Date", "Customer GSTIN", "Taxable", "CGST", "SGST", "Total GST", "Grand Total"]
        widths = [18, 14, 22, 14, 12, 12, 14, 14]
        for c, (h, w) in enumerate(zip(heads, widths), 1):
            _header_cell(ws1, 2, c, h, w)

        for r, s in enumerate(sales, 3):
            for c, v in enumerate([
                s.invoice_no, s.sale_date.strftime("%d-%m-%Y"),
                s.customer.gstin if s.customer and s.customer.gstin else "URP",
                s.taxable_amount, s.cgst_amount, s.sgst_amount,
                s.total_gst, s.grand_total,
            ], 1):
                _data_cell(ws1, r, c, v)

        ws2 = wb.create_sheet("Inward Supplies")
        ws2.merge_cells("A1:H1")
        ws2["A1"].value = f"Inward Supplies (Purchases) – {period}"
        ws2["A1"].font = Font(bold=True, size=12, color="1E3A5F")

        for c, (h, w) in enumerate(zip(["GRN No.", "Date", "Vendor GSTIN", "Taxable", "CGST", "SGST", "ITC", "Grand"], widths), 1):
            _header_cell(ws2, 2, c, h, w)

        for r, p in enumerate(purchases, 3):
            for c, v in enumerate([
                p.grn_no, p.purchase_date.strftime("%d-%m-%Y"),
                p.vendor.gstin if p.vendor and p.vendor.gstin else "—",
                p.taxable_amount, p.cgst_amount, p.sgst_amount,
                p.total_gst, p.grand_total,
            ], 1):
                _data_cell(ws2, r, c, v)

        path = _out_path("gst_report")
        wb.save(path)
        return path
    finally:
        session.close()